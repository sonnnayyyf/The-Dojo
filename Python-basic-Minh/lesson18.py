import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import pandas as pd
1.
# def load_june(filename):
#     transactions = []
#     skip = 0
#     with open(filename) as f:
#         reader = csv.DictReader(f)
#         for i in reader:
#             try:
#                 datetime.strptime(i['date'], "%Y-%m-%d")
#                 i['amount'] = float(i['amount'])
#                 transactions.append(i)
#             except (ValueError, KeyError):
#                 skip += 1
#     print(f"✅  Loaded {len(transactions)} valid rows. Skipped {skip} junk rows ")
#     return transactions

# load_june('june_transactions.csv')


2

# def summary(rows):
#     income = 0
#     expenses = 0
#     for i in rows:
#         if i['category'] == 'Income':
#             income += i['amount']
#         else:
#             expenses += i['amount']
#     net = income - expenses
#     trans_dict = {
#         'income': income,
#         'expenses': expenses,
#         'net': net
#     }
#     print(f'✅  Income: {income:.2f} | Expenses: {expenses:.2f} | Net: {net:.2f}')
#     return trans_dict

# summary(load_june('june_transactions.csv'))

3.
# def biggest_expense_day(rows):
#     highest_amount = 0
#     for i in rows:
#         if i['category'] != 'Income':
#             if i['amount'] > highest_amount:
#                 highest_amount = i['amount']
#     for j in rows:
#         if j['amount'] == highest_amount:
#             highest_amount_date = j['date']
#     print(f'Biggest spending day: {highest_amount_date} — ${highest_amount}')

# biggest_expense_day(load_june('june_transactions.csv'))

4.

# def category_report(rows):
#     new_dict = {}
#     total_expenses = 0
#     for i in rows:
#         if i['category'] == 'Income':
#             continue
#         if i['category'] not in new_dict:
#             new_dict[i['category']] = 0
#         new_dict[i['category']] += i['amount']
#         total_expenses += i['amount']
#     for cate,amt in new_dict.items():
#         percentage = amt/total_expenses * 100
#         print(f'{cate}             ${amt:.2f}   {percentage:.1f}%')

        
        



# category_report(load_june('june_transactions.csv'))

5.
# def flag_large(rows, threshold = 500):
#     trans_list = []
#     count = 0
#     for i in rows:
#         if i['amount'] > threshold:
#             trans_list.append(i)
#             count += 1
#     trans_list.sort(key=lambda i: -i['amount'])
#     print(trans_list)
#     for j in trans_list:
#         print(f'{j["date"]} | {j['category']} | {j['description']} | {j['amount']}')
#     print(f'✅ There are {count} transactions over $500 in this file.')            

# flag_large(load_june('june_transactions.csv'), threshold = 500)

6.

# def weekly_expenses(rows):
#     transaction_week = {
#         '1': 0.0,
#         '2': 0.0,
#         '3': 0.0,
#         '4': 0.0
#     }

#     for i in rows:
#         if i['category'] == "Income":
#             continue
#         date = datetime.strptime(i['date'], "%Y-%m-%d")
#         if date.day <= 7:
#             transaction_week['1'] += i['amount']
#         elif 7< date.day <= 14:
#             transaction_week['2'] += i['amount']
#         elif 14< date.day <= 21:
#             transaction_week['3'] += i['amount']
#         elif 21< date.day <= 31:
#             transaction_week['4'] += i['amount']
#     for day, amt in transaction_week.items():
#         print(f'Week {day} : ${amt:.2f}')
#     print(f'✅ All four weeks should have a non-zero total. Weeks sum to ${sum(transaction_week.values())}.')


# weekly_expenses(load_june('june_transactions.csv'))


7.


# def B1(filename):
#     df = pd.read_csv(filename)
#     df['amount'] = pd.to_numeric(df['amount'], errors = 'coerce')
#     df['date'] = pd.to_datetime(df['date'], errors = 'coerce', format = 'mixed')
#     df = df.dropna(subset=['amount', 'date'])
#     expenses = df[df['category'] != 'Income']['amount'].sum()
#     income = df[df['category'] == 'Income']['amount'].sum()
#     print(f"[panda] Income: {income:.2f} | expenses: {expenses:.2f} | net: {(income - expenses):.2f}")
# print(B1('june_transactions.csv'))


8.


    
# def B2(filename):
#     df = pd.read_csv(filename)
#     df['amount'] = pd.to_numeric(df['amount'], errors = 'coerce')
#     df['date'] = pd.to_datetime(df['date'], errors = 'coerce', format = 'mixed')
#     df = df.dropna(subset=['amount', 'date'])
#     expenses_only = df[df['category'] != 'Income']
#     expense_separated = expenses_only.groupby('category')['amount'].sum().reset_index()
#     expense_separated.to_excel('june_report.xlsx', index = False, startrow = 1)
#     wb = load_workbook('june_report.xlsx')
#     ws = wb.active
#     ws.append(['Total expense', '=SUM(B3:B9)'])
#     ws['A1'] = "June 2025 — Expense Report"
#     ws.merge_cells('A1:B1')
#     ws.column_dimensions['A'].width = 20
#     ws.column_dimensions['B'].width = 20
#     ws['A1'].font = Font(bold = True, size = 16)
#     ws['A1'].alignment = Alignment (horizontal = 'center')
#     ws['B2'].alignment = Alignment (horizontal = 'right')
#     ws['A1'].fill = PatternFill('solid',start_color = '4472C4' )
#     ws['A10'].font = Font(bold = True, size = 14)
#     ws['B10'].font = Font(bold = True, size = 14)
#     for grid in range(3,11):
#         ws.cell(row=grid, column = 2).number_format = "$#,##0.00"
#         ws.cell(row=grid, column = 2).fill = PatternFill('solid', start_color = '66FF33')
#         ws.cell(row=grid, column = 1).fill = PatternFill('solid', start_color = '66FF33')
#     ws['A2'].font = Font(bold = True, size = 12)
#     ws['B2'].font = Font(bold = True, size = 12)
#     ws['A2'].fill = PatternFill('solid', start_color = 'DD61CE')
#     ws['B2'].fill = PatternFill('solid', start_color = 'DD61CE')

#     wb.save('june_report.xlsx')



# B2('june_transactions.csv')


9

# def top_vendors(filename, n=3):
#     transactions_dict = {}
#     df = pd.read_csv(filename)
#     df['amount'] = pd.to_numeric(df['amount'], errors = 'coerce')
#     df['date'] = pd.to_datetime(df['date'], errors = 'coerce', format = 'mixed')
#     df = df.dropna(subset=['amount', 'date'])
#     expense_only = df[df['category'] != 'Income']
#     descriptions_frequency = expense_only['description'].value_counts()
#     print(descriptions_frequency.head(n))

        

    

# top_vendors('june_transactions.csv', n=3)



10

def running_balance(filename):
    df = pd.read_csv(filename)
    df['amount'] = pd.to_numeric(df['amount'], errors = 'coerce')
    df['date'] = pd.to_datetime(df['date'], errors = 'coerce', format = 'mixed')
    df = df.dropna(subset=['amount', 'date'])
    df = df.sort_values('date')
    expense_max = df[df['category'] != 'Income']['amount'].max()
    balance = 0
    for index, i in df.iterrows():
        if i['category'] == 'Income':
            balance += i['amount']
            sign = '+'
        else:
            balance -= i['amount']
            sign = '-'
        print(f'{i['date'].date()} | {i['description']} | {sign}${i['amount']:.2f} | \nBalance: {balance:.2f}')
    for index, j in df.iterrows():
        if j['amount'] == expense_max:
            print(f'{j['date']} | {j['description']} | -${j['amount']}')
            



running_balance('june_transactions.csv')

