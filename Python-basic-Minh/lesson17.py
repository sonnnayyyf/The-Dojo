from openpyxl import load_workbook, Workbook

from openpyxl.styles import Font, PatternFill, Alignment

# wb = load_workbook('transactions.xlsx')

# print(wb.sheetnames)
# ws = wb.active
# print(ws.max_row, ws.max_column)


# print(ws['B1'].value)
# print(ws.cell(row=3,column =4).value)

# for row in ws.iter_rows(min_row = 2, values_only = True):
#     print(row)

# for row in ws.iter_rows(min_row=2, values_only=True):
#     date, description, category, amount = row
#     print(category, amount)

# wb = Workbook()
# ws = wb.active           # its first sheet
# ws.title = 'Report'   

# ws['A1'] = 'My Report'

# ws.append(['Category', 'Total'])
# ws.append(['Groceries', 290.95])



# wb.save('my_report.xlsx')

# wb = load_workbook('my_report.xlsx')
# ws = wb.active
# # ws.append(['Insurance', 390.67])
# # ws.append(['Equipment', 400.32])
# # ws.append(['Total', '=SUM(B3:B5)'])

# ws['A2'].font = Font(bold = True, size = 14, color = '00FF00')
# ws['A3'].fill = PatternFill('solid', start_color='4472C4')

# ws['A1'].alignment = Alignment (horizontal = 'center')
# ws.column_dimensions['A'].width = 18


# ws['B3'].number_format = '$#,##0.00'
# ws['C2'].number_format = '0.0%'



# wb.save('my_report.xlsx')

1.

# wb = load_workbook('transactions.xlsx')

# print(wb.sheetnames)
# ws = wb.active
# print(ws.max_row, ws.max_column )

2.

# wb = load_workbook('transactions.xlsx')
# ws = wb.active
# print(ws['B2'].value)
# print(ws['D2'].value)

# print(ws.cell(row = 2, column = 2).value)
# print(ws.cell(row = 2, column = 4).value)

3.
# wb = load_workbook('transactions.xlsx')
# ws = wb.active
# total_amount = 0
# for row in ws.iter_rows(min_row = 2, values_only = True):
#     date, description, category, amount = row
    
#     if category != "Income":
#         total_amount += amount
# print(total_amount)
    

4.

# wb = Workbook()
# ws = wb.active
# ws.title = 'Budget'
# ws['A1'] = 'Category'
# ws['B1'] = 'Limit'
# ws.append(['Groceries', 400])
# ws.append(['Transport', 150])
# ws.append(['Food', 200])
# wb.save('budget.xlsx')



# instruction
# wb = Workbook()
# ws = wb.active           # its first sheet
# ws.title = 'Report'   

# ws['A1'] = 'My Report'

# ws.append(['Category', 'Total'])
# ws.append(['Groceries', 290.95])

5.
# wb1 = load_workbook('transactions.xlsx')
# ws1 = wb1.active

# wb2 = Workbook()
# ws2 = wb2.active
# ws2.title = 'expenses'

# instruction
# wb = Workbook()
# ws = wb.active           # its first sheet
# ws.title = 'Report'   

# ws['A1'] = 'My Report'

# ws2.append(['Date', 'Description', 'Category', 'Amount'])

# for row in ws1.iter_rows(min_row = 2, values_only = True):
#     date, description, category, amount = row
#     if category != "Income":
#         ws2.append(row)
# wb2.save('expenses.xlsx')

6.
# instruction
# wb = load_workbook('my_report.xlsx')
# ws = wb.active
# # ws.append(['Insurance', 390.67])
# # ws.append(['Equipment', 400.32])
# # ws.append(['Total', '=SUM(B3:B5)'])

# ws['A2'].font = Font(bold = True, size = 14, color = '00FF00')
# ws['A3'].fill = PatternFill('solid', start_color='4472C4')

# ws['A1'].alignment = Alignment (horizontal = 'center')
# ws.column_dimensions['A'].width = 18


# ws['B3'].number_format = '$#,##0.00'
# ws['C2'].number_format = '0.0%'




# wb = load_workbook('budget.xlsx')
# ws = wb.active
# for cell in ws[1]:
#     cell.font = Font(bold = True, color = 'FFFFFF')
#     cell.fill = PatternFill('solid', start_color = '4472C4')
#     cell.alignment = Alignment (horizontal = 'center')
# ws.column_dimensions['A'].width = 18
# ws.column_dimensions['B'].width = 18


# ws.append(['Total', '=SUM(B2:B4)'])
# for grid in range(2,6):
#     ws.cell(row = grid, column = 2).number_format = "$#,##0.00"
# wb.save('budget.xlsx')



# stretch goal

import pandas as pd

df = pd.read_excel('transactions.xlsx')


expenses_only = df[df['Category'] != 'Income']
expense_excel = expenses_only.groupby('Category')['Amount'].sum().reset_index()
expense_excel.columns = ['Category', 'Total']
expense_excel['Total'] = expense_excel['Total'].round(2)

expense_excel.to_excel('monthly_report.xlsx', index = False)

wb = load_workbook('monthly_report.xlsx')
ws = wb.active



for cell in ws[1]:
    cell.font = Font(bold = True)
    cell.fill = PatternFill('solid', start_color = '4472C4')
    cell.alignment = Alignment (horizontal = 'center')
ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 18

ws.append(['Total', '=SUM(B2:B8)'])

for grid in range(2,9):
    ws.cell(row = grid, column = 2).number_format = "$#,##0.00"


ws.append(['Expense Report - January 2025'])
ws['A10'].font = Font(bold = True)
ws.merge_cells('A10:B10')
for cell in ws[10]:
    cell.alignment = Alignment (horizontal = 'center')

wb.save('monthly_report.xlsx')
