import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import pandas as pd

1
def load_invoices(filepath):
    invoice_dict = {}
    with open(filepath) as f:
        reader = csv.DictReader(f)
        for i in reader:
            try:
                i['amount_due'] = float(i['amount_due'])
                i['due_date'] = datetime.strptime(i['due_date'],"%Y-%m-%d" )
            except (ValueError, KeyError):
                continue
            invoice_dict[i['invoice_id']] = i
    return invoice_dict
    



# print(load_invoices('invoices.csv'))
2

def total_payments(filepath):
    total_payment_dict = {}
    latest_date_dict = {}
    with open(filepath) as f:
        reader = csv.DictReader(f)
        for i in reader:
            try:
                i['amount_paid'] = float(i['amount_paid'])
                i['date_paid'] = datetime.strptime(i['date_paid'], "%Y-%m-%d")
            except (ValueError, KeyError):
                continue

            if i['invoice_id'] not in total_payment_dict:
                total_payment_dict[i['invoice_id']] = i['amount_paid']
            else:
                total_payment_dict[i['invoice_id']] += i['amount_paid']
            
            if i['invoice_id'] not in latest_date_dict or i['date_paid'] > latest_date_dict[i['invoice_id']]:
                latest_date_dict[i['invoice_id']] = i['date_paid']
    return total_payment_dict, latest_date_dict

# print(total_payments('payments.csv'))

3
def classficiation(invoice_dict, total_payment_dict, latest_date_dict):
    invoice_type_dict = {
        "Unpaid": [],
        "Paid in full": [],
        "Partially paid": [],
        "Overpaid": [],
        "Late": []
    }
    for i in invoice_dict:
        invoice = invoice_dict[i]
        amount = invoice['amount_due']
        total_partially_paid_invoice = 0
        overpaid_invoice = 0
        pay = total_payment_dict.get(i, 0)
        if pay == 0:
            invoice_type_dict["Unpaid"].append(i)
        elif abs(pay - amount) < 0.01:
            invoice_type_dict["Paid in full"].append(i)
        elif pay > 0 and pay < amount:
            invoice_type_dict["Partially paid"].append(i)
            total_partially_paid_invoice += amount
        else:
            invoice_type_dict["Overpaid"].append(i)
            overpaid_invoice += amount
        
        due = invoice['due_date']
        if i in latest_date_dict:
            payment_date = latest_date_dict[i]
            if payment_date > due:
                invoice_type_dict['Late'].append(i)

    for a in invoice_type_dict:
        print(f'{a} | {len(invoice_type_dict[a])}')
    return invoice_type_dict, overpaid_invoice, total_partially_paid_invoice
        

totals, latests = total_payments('payments.csv')

classficiation(load_invoices('invoices.csv'), totals, latests)
invoice_type_dict, overpaid, partially_paid = classficiation(load_invoices('invoices.csv'), totals, latests)

4
def invoice_reconciliation(invoice_dict, total_payment_dict, invoice_type_dict, overpaid, partially_paid ):
    total_invoice = 0
    total_received = 0
    for i in invoice_dict:
        invoice = invoice_dict[i]
        amount = invoice['amount_due']
        total_invoice += amount
    for j in total_payment_dict:
        total_received += total_payment_dict[j]
    outstanding = total_invoice - total_received
    print("=== INVOICE RECONCILIATION ===")
    print(f'PAID IN FULL ({len(invoice_type_dict["Paid in full"])}):')
    print("  ", end="")
    for a in invoice_type_dict['Paid in full']: 
        print(f'{a}', end=", ")
    print()
    print()
    
    print(f'PARTIALLY PAID ({len(invoice_type_dict["Partially paid"])}):')
    for a in invoice_type_dict['Partially paid']:
        invoice = invoice_dict[a]
        amount = invoice["amount_due"]
        paid = total_payment_dict.get(a)
        print(f'  {a} - short by ${(amount - paid):.2f}')
    print()

    print(f'UNPAID ({len(invoice_type_dict["Unpaid"])}):')
    print("  ", end="")
    for a in invoice_type_dict['Unpaid']: 
        print(f'{a}', end=", ")
    print()
    print()

    print(f'OVERPAID ({len(invoice_type_dict["Overpaid"])}):')
    print("  ", end="")
    for a in invoice_type_dict['Overpaid']: 
        invoice = invoice_dict[a]
        amount = invoice["amount_due"]
        paid = total_payment_dict.get(a)
        print(f'  {a} - over by ${(paid - amount):.2f}')
    print()

    print(f'LATE PAYMENTS ({len(invoice_type_dict["Late"])}):')
    print("  ", end="")
    for a in invoice_type_dict['Late']: 
        print(f'{a}', end=", ")
    print()
    print()
    
    print(f'Total invoiced: ${total_invoice}')
    print(f'Total received: ${total_received}')
    print(f'Outstanding: ${outstanding}')
          


invoice_reconciliation(load_invoices('invoices.csv'), totals, invoice_type_dict, overpaid, partially_paid)




5.

def reconciliation_report(invoice_dict,total_payment_dict, latest_date_dict,invoice_type_dict):
    wb = load_workbook('reconciliation_report.xlsx')
    ws = wb.active
    ws.title = "Invoice Reconciliation Report"
    ws.append(['Invoice Reconciliation Report'])
    headers = ['Invoice ID', 'Client', 'Amount Due', 'Amount Paid', 'Status', 'Days Late']
    ws.append(headers)
    ws['A1'].alignment = Alignment (horizontal = 'center')
    ws['A1'].font = Font(bold = True, size = 20)
    ws.merge_cells('A1:F1')

    for cell in ws[2]:
        cell.font = Font(bold =True)   
        cell.fill = PatternFill('solid',start_color = '4472C4' )
    
    for i in invoice_dict:
        invoice = invoice_dict[i]
        paid = total_payment_dict.get(i, 0)
        due = invoice["amount_due"]
        due_date = invoice['due_date']
        days_late = 0
        
        if i in latest_date_dict:
            date_paid = latest_date_dict[i]
            if date_paid > due_date:
                days_late = (date_paid - due_date).days



        
        if paid == 0:
            status = 'Unpaid'
        elif abs(paid - due) < 0.01:
            status = 'Paid in full'
        elif paid > due:
            status = 'Overpaid'
        elif paid > 0 and paid < due:
            status = 'Partially paid'
        

        ws.append([
            i,
            invoice["client"],
            due,
            paid,
            status,
            days_late

        ])


    
    ws.append([' ', 'Total', '=SUM(C3:C14)', '=SUM(D3:D14)'])
    col = [ 'C', 'D']
    for a in col:
        ws.column_dimensions[a].width = 12

        
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['B'].width = 18

    for grid in range(3,16):
        ws.cell(row = grid, column = 3).number_format = "$#,##0.00"
        ws.cell(row = grid, column = 4).number_format = "$#,##0.00"
    
    wb.save('reconciliation_report.xlsx')


    

reconciliation_report(load_invoices('invoices.csv'), totals, latests, invoice_type_dict)


    
    

# Row:    
# for cell in ws[1]:
#     cell.font =

# Column:
# for cell in range(2, 6):
#     cell.font =



# Stretch goal

def function(invoice_type_dict, invoice_dict):
    wb = load_workbook('reconciliation_report.xlsx')
    summary = wb.create_sheet('Summary')
    header = ['Status', 'Count', 'Total Amount Due']
    summary.append(header)
  
    for status in ['Unpaid', 'Paid in full', 'Partially paid', 'Overpaid']:
        total = 0
        for i in invoice_type_dict[status]:
            invoice = invoice_dict[i]
            total += invoice['amount_due']
        
        summary.append([
            status,
            len(invoice_type_dict[status]),
            total
        ])

    wb.save('reconciliation_report.xlsx')

function(invoice_type_dict, load_invoices('invoices.csv'))
